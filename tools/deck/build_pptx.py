#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Kairos Deck v2 -> 原生可编辑 PPTX（16:9，黑底 + 币安金）。读 slides.py。
用 /usr/bin/python3 运行（已装 python-pptx 1.0.2）。3D 图作图片置入；文字/表可编辑。
webp 经 sips 转 png。
"""
import os, sys, subprocess
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
from slides import SLIDES, LOGO_SM  # noqa
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUT = os.path.join(ROOT, "output", "deck", "Kairos-Deck-zh.pptx")
CACHE = os.path.join(ROOT, "output", "deck-assets", "_pptx")
GOLD = RGBColor(0xF0, 0xB9, 0x0B); GOLD2 = RGBColor(0xFC, 0xD5, 0x35); GOLDD = RGBColor(0xB8, 0x86, 0x0B)
CYAN = RGBColor(0x16, 0xD2, 0xE8); WHITE = RGBColor(0xFF, 0xFF, 0xFF); INK = RGBColor(0x0B, 0x0B, 0x0B)
BODY = RGBColor(0xC2, 0xC2, 0xC2); MUT = RGBColor(0x8A, 0x8A, 0x8A); CARD = RGBColor(0x16, 0x16, 0x16)
LINE = RGBColor(0x33, 0x33, 0x33); F = "PingFang SC"
SW, SH = Inches(13.333), Inches(7.5); MX = Inches(0.7)


def png_for(rel):
    p = os.path.join(ROOT, rel)
    if not os.path.isfile(p):
        return None
    if p.lower().endswith((".png", ".jpg", ".jpeg")):
        return p
    os.makedirs(CACHE, exist_ok=True)
    out = os.path.join(CACHE, os.path.splitext(os.path.basename(p))[0] + ".png")
    if not os.path.isfile(out):
        subprocess.run(["sips", "-s", "format", "png", p, "--out", out], capture_output=True)
    return out if os.path.isfile(out) else None


def _ea(run, font):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", font)


def text(sl, s, l, t, w, h, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp=None, font=F):
    tb = sl.shapes.add_textbox(l, t, w, h); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    first = True
    for line in str(s).split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align
        if sp:
            p.line_spacing = sp
        r = p.add_run(); r.text = line; f = r.font
        f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = font; _ea(r, font)
    return tb


def rect(sl, l, t, w, h, fill=CARD, line=LINE, lw=1.0, rad=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = sl.shapes.add_shape(shape, l, t, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    try:
        sp.adjustments[0] = rad
    except Exception:
        pass
    return sp


def fit(sl, rel, l, t, w, h):
    p = png_for(rel)
    if not p:
        return None
    pic = sl.shapes.add_picture(p, l, t, width=w)
    if pic.height > h:
        ratio = pic.width / pic.height; pic.height = h; pic.width = int(h * ratio)
    pic.left = int(l + (w - pic.width) / 2); pic.top = int(t + (h - pic.height) / 2)
    return pic


def bg(sl):
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = INK


def head(sl, s):
    if s.get("kicker"):
        text(sl, s["kicker"], MX, Inches(0.5), Inches(8), Inches(0.3), 12, GOLD, bold=True)
    rect(sl, MX, Inches(0.86), Inches(0.7), Pt(3), fill=GOLD, line=None, rad=0)
    if s.get("no"):
        text(sl, f'{s["no"]}', Inches(11.6), Inches(0.5), Inches(1.1), Inches(0.3), 14, MUT, align=PP_ALIGN.RIGHT)
    text(sl, s["title"], MX, Inches(1.0), Inches(12), Inches(1.0), 34, WHITE, bold=True, sp=1.04)
    y = Inches(1.65) + Inches(0.4) * (str(s["title"]).count("\n"))
    if s.get("lede"):
        text(sl, s["lede"], MX, y, Inches(11.6), Inches(0.9), 15, BODY, sp=1.25)
        y = y + Inches(0.85)
    return y


def brand(sl):
    p = png_for(LOGO_SM)
    if p:
        sl.shapes.add_picture(p, Inches(11.95), Inches(6.98), height=Inches(0.3))
    text(sl, "KAIROS", Inches(12.3), Inches(6.99), Inches(0.9), Inches(0.3), 10, GOLD, bold=True)


# ---------------- layouts ----------------
def L_cover(sl, s):
    p = png_for(s["hero"])
    if p:
        pic = sl.shapes.add_picture(p, Inches(4.6), 0, height=SH)  # right-fill
        pic.left = SW - pic.width
    rect(sl, 0, 0, Inches(6.6), SH, fill=INK, line=None, rad=0)  # left scrim
    lp = png_for(LOGO_SM)
    if lp:
        sl.shapes.add_picture(lp, MX, Inches(0.55), height=Inches(0.42))
    text(sl, "Kairos", MX + Inches(0.55), Inches(0.55), Inches(3), Inches(0.45), 20, GOLD, bold=True)
    text(sl, s["kicker"], MX, Inches(2.45), Inches(6), Inches(0.3), 12, GOLD, bold=True)
    text(sl, s["brand"], MX, Inches(2.8), Inches(6.5), Inches(1.4), 84, GOLD, bold=True)
    text(sl, s["sub"], MX, Inches(4.4), Inches(5.6), Inches(0.6), 19, GOLD2, bold=True, sp=1.1)
    text(sl, s["lede"], MX, Inches(5.15), Inches(5.4), Inches(0.7), 15, BODY, sp=1.3)
    text(sl, "◆ " + s["chain"], MX, Inches(5.95), Inches(6), Inches(0.3), 12, BODY, bold=True)
    text(sl, "BACKED BY WORLD-CLASS INVESTORS", MX, Inches(6.5), Inches(6), Inches(0.25), 9, MUT, bold=True)
    x = MX
    for fimg in s["backed_by"]:
        tw = Inches(0.78)
        rect(sl, x, Inches(6.85), tw, Inches(0.4), fill=WHITE, line=None, rad=0.12)
        fit(sl, fimg, x + Inches(0.07), Inches(6.9), tw - Inches(0.14), Inches(0.3))
        x += tw + Inches(0.1)


def L_investors(sl, s):
    head(sl, s)
    vcs = s["vcs"]; tw, th, gap = Inches(1.92), Inches(0.92), Inches(0.13); top = Inches(2.75)
    for i, v in enumerate(vcs):
        r, c = divmod(i, 6); l = MX + c * (tw + gap); t = top + r * (th + gap)
        rect(sl, l, t, tw, th, fill=WHITE, line=None, rad=0.1)
        fit(sl, v["file"], l + Inches(0.16), t + Inches(0.13), tw - Inches(0.32), th - Inches(0.26))
    caps = [v for v in vcs if v.get("cap")]; cy = Inches(4.75)
    for i, v in enumerate(caps):
        col = i % 2; l = MX + col * Inches(6.1); t = cy + (i // 2) * Inches(0.4)
        text(sl, f'{v["name"]} — {v["cap"]}', l, t, Inches(5.9), Inches(0.38), 12, BODY)


def L_cards(sl, s):
    y = head(sl, s); cols = s.get("cols", 2); cards = s["cards"]
    gap = Inches(0.25); rows = (len(cards) + cols - 1) // cols
    cw = int((SW - 2 * MX - gap * (cols - 1)) / cols)
    ch = int((Inches(6.7) - y - gap * (rows - 1)) / rows)
    for i, c in enumerate(cards):
        r, cc = divmod(i, cols); l = MX + cc * (cw + gap); t = y + r * (ch + gap)
        rect(sl, l, t, cw, ch)
        rect(sl, l, t, cw, Pt(3), fill=GOLD, line=None, rad=0)
        yy = t + Inches(0.22)
        if c.get("tag"):
            text(sl, c["tag"], l + Inches(0.3), yy, cw - Inches(0.6), Inches(0.25), 10, GOLD, bold=True)
            yy += Inches(0.3)
        text(sl, c["h"], l + Inches(0.3), yy, cw - Inches(0.6), Inches(0.5), 20, GOLD, bold=True)
        text(sl, c["t"], l + Inches(0.3), yy + Inches(0.5), cw - Inches(0.6), ch - Inches(1.2), 14, BODY, sp=1.25)


def L_statgrid(sl, s):
    head(sl, s); stats = s["stats"]; gap = Inches(0.25); top = Inches(2.5)
    cw = int((SW - 2 * MX - gap * 2) / 3); ch = Inches(1.7)
    for i, st in enumerate(stats):
        r, c = divmod(i, 3); l = MX + c * (cw + gap); t = top + r * (ch + gap)
        rect(sl, l, t, cw, ch)
        text(sl, st["lab"], l + Inches(0.3), t + Inches(0.24), cw - Inches(0.6), Inches(0.5), 13, MUT)
        sz = 50 if st.get("hero") else (26 if st.get("small") else 40)
        col = GOLD if st.get("hero") else WHITE
        text(sl, st["num"], l + Inches(0.3), t + Inches(0.78), cw - Inches(0.6), Inches(0.8), sz, col, bold=True)
    if s.get("note"):
        text(sl, s["note"], MX, Inches(6.65), Inches(12), Inches(0.3), 11, RGBColor(0x5C, 0x5C, 0x5C))


def L_concept(sl, s):
    head(sl, s); Lw = Inches(6.5) if (s.get("hero") or s.get("barviz")) else Inches(11.8); y = Inches(2.9)
    if s.get("callout"):
        rect(sl, MX, y, Lw, Inches(1.45), fill=RGBColor(0x1A, 0x15, 0x05), line=GOLD)
        text(sl, s["callout"], MX + Inches(0.3), y + Inches(0.22), Lw - Inches(0.6), Inches(0.6), 30, GOLD, bold=True)
        text(sl, s.get("callout_sub", ""), MX + Inches(0.3), y + Inches(0.9), Lw - Inches(0.6), Inches(0.4), 13, BODY)
        y += Inches(1.65)
    if s.get("example"):
        text(sl, s["example"], MX, y, Lw, Inches(0.8), 16, RGBColor(0xE6, 0xE6, 0xE6), sp=1.3)
    if s.get("chips"):
        x, cy = MX, Inches(4.95)
        for ch in s["chips"]:
            w = Inches(0.5 + 0.2 * len(ch))
            if x + w > MX + Inches(6.5):
                x = MX; cy += Inches(0.55)
            rect(sl, x, cy, w, Inches(0.42), fill=RGBColor(0x14, 0x14, 0x14), line=LINE, rad=0.5)
            text(sl, ch, x, cy + Inches(0.07), w, Inches(0.3), 13, RGBColor(0xE6, 0xE6, 0xE6), align=PP_ALIGN.CENTER)
            x += w + Inches(0.16)
    if s.get("hero"):
        fit(sl, s["hero"], Inches(7.5), Inches(2.5), Inches(5.2), Inches(4.4))
    elif s.get("barviz"):
        bx, bw = Inches(7.6), Inches(5.0)
        text(sl, "＝ $1.00", bx, Inches(2.9), bw, Inches(0.5), 26, GOLD, bold=True, align=PP_ALIGN.CENTER)
        rect(sl, bx, Inches(3.6), int(bw * 0.18), Inches(1.1), fill=GOLD2, line=None, rad=0.04)
        rect(sl, bx + int(bw * 0.18), Inches(3.6), int(bw * 0.82), Inches(1.1), fill=CYAN, line=None, rad=0.04)
        text(sl, "YES 0.18", bx, Inches(3.95), int(bw * 0.18), Inches(0.4), 11, INK, bold=True, align=PP_ALIGN.CENTER)
        text(sl, "NO  0.82", bx + int(bw * 0.18), Inches(3.95), int(bw * 0.82), Inches(0.4), 16, INK, bold=True, align=PP_ALIGN.CENTER)
        text(sl, "买 YES @0.18", bx, Inches(4.85), Inches(2.2), Inches(0.4), 13, GOLD, bold=True)
        text(sl, "≡ 卖 NO @0.82", bx + Inches(2.9), Inches(4.85), Inches(2.1), Inches(0.4), 13, CYAN, bold=True, align=PP_ALIGN.RIGHT)


def L_cols(sl, s):
    y = head(sl, s); items = s["items"]; cols = s.get("cols", 3)
    gap = Inches(0.3); rows = (len(items) + cols - 1) // cols
    cw = int((SW - 2 * MX - gap * (cols - 1)) / cols); ch = int((Inches(6.6) - y - gap * (rows - 1)) / max(rows, 1))
    for i, it in enumerate(items):
        r, c = divmod(i, cols); l = MX + c * (cw + gap); t = y + r * (ch + gap)
        text(sl, it["n"], l, t, cw, Inches(0.35), 15, GOLD, bold=True)
        text(sl, it["h"], l, t + Inches(0.4), cw, Inches(0.5), 21, WHITE, bold=True)
        text(sl, it["t"], l, t + Inches(0.95), cw, ch - Inches(1.0), 14.5, BODY, sp=1.3)


def L_flow(sl, s):
    head(sl, s); steps = s["steps"]; n = len(steps)
    gap = Inches(0.4); cw = int((SW - 2 * MX - gap * (n - 1)) / n); t = Inches(3.0); ch = Inches(2.7)
    for i, st in enumerate(steps):
        l = MX + i * (cw + gap)
        rect(sl, l, t, cw, ch)
        cx = l + cw / 2
        rect(sl, int(cx - Inches(0.36)), t + Inches(0.3), Inches(0.72), Inches(0.72), fill=GOLD, line=None, rad=0.5, shape=MSO_SHAPE.OVAL)
        text(sl, st["b"], int(cx - Inches(0.36)), t + Inches(0.45), Inches(0.72), Inches(0.45), 20, INK, bold=True, align=PP_ALIGN.CENTER)
        text(sl, st["h"], l + Inches(0.3), t + Inches(1.3), cw - Inches(0.6), Inches(0.5), 21, WHITE, bold=True, align=PP_ALIGN.CENTER)
        text(sl, st["t"], l + Inches(0.3), t + Inches(1.85), cw - Inches(0.6), Inches(0.8), 14, BODY, sp=1.25, align=PP_ALIGN.CENTER)
        if i < n - 1:
            text(sl, "→", l + cw, t + Inches(0.5), gap, Inches(0.6), 26, GOLD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def L_spec(sl, s):
    head(sl, s); rows = [s["headers"]] + s["rows"]; nr, nc = len(rows), 3
    t = Inches(2.55); tbl = sl.shapes.add_table(nr, nc, MX, t, SW - 2 * MX, Inches(4.0)).table
    tbl.columns[0].width = Inches(2.6); tbl.columns[1].width = Inches(4.6); tbl.columns[2].width = Inches(4.73)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Inches(0.66)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci); cell.margin_left = Inches(0.24); cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = (GOLD if ci == 2 else RGBColor(0x14, 0x14, 0x14))
                col, bold = (INK if ci == 2 else RGBColor(0xCF, 0xCF, 0xCF)), True
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0x10, 0x10, 0x10)
                col, bold = (GOLD, True) if ci == 0 else ((WHITE, True) if ci == 2 else (MUT, False))
            p = cell.text_frame.paragraphs[0]; r = p.add_run(); r.text = str(val)
            r.font.size = Pt(15); r.font.bold = bold; r.font.color.rgb = col; r.font.name = F; _ea(r, F)


def L_fcards(sl, s):
    head(sl, s); cards = s["cards"]; n = len(cards); gap = Inches(0.25)
    cw = int((SW - 2 * MX - gap * (n - 1)) / n); t = Inches(2.7); ch = Inches(3.3)
    for i, c in enumerate(cards):
        l = MX + i * (cw + gap); rect(sl, l, t, cw, ch)
        fit(sl, c["icon"], l + (cw - Inches(1.3)) / 2, t + Inches(0.25), Inches(1.3), Inches(1.3))
        text(sl, c["tag"], l, t + Inches(1.6), cw, Inches(0.25), 10, GOLD, bold=True, align=PP_ALIGN.CENTER)
        text(sl, c["h"], l, t + Inches(1.9), cw, Inches(0.4), 18, WHITE, bold=True, align=PP_ALIGN.CENTER)
        text(sl, c["t"], l + Inches(0.2), t + Inches(2.4), cw - Inches(0.4), Inches(0.8), 13, BODY, sp=1.3, align=PP_ALIGN.CENTER)


def L_eco(sl, s):
    head(sl, s); cx, cy = Inches(6.66), Inches(4.7)
    for rx, ry in ((Inches(1.7), Inches(1.25)), (Inches(3.1), Inches(1.7)), (Inches(4.6), Inches(2.0))):
        rect(sl, int(cx - rx), int(cy - ry), int(2 * rx), int(2 * ry), fill=None, line=LINE, lw=0.75, rad=0, shape=MSO_SHAPE.OVAL)
    rect(sl, int(cx - Inches(0.6)), int(cy - Inches(0.6)), Inches(1.2), Inches(1.2), fill=RGBColor(0x2A, 0x20, 0x05), line=GOLD, rad=0.5, shape=MSO_SHAPE.OVAL)
    text(sl, s["center"], int(cx - Inches(0.6)), int(cy - Inches(0.35)), Inches(1.2), Inches(0.7), 13, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    inner_pos = [(0, -2.0), (3.4, -1.0), (3.4, 1.0), (0, 2.1), (-3.4, 1.0), (-3.4, -1.0)]
    for (dx, dy), label in zip(inner_pos, s["inner"]):
        l = int(cx + Inches(dx) - Inches(0.9)); t = int(cy + Inches(dy) - Inches(0.18))
        text(sl, "●  " + label, l, t, Inches(2.0), Inches(0.36), 12.5, BODY, align=PP_ALIGN.CENTER)
    out_pos = [(0, -2.55, "Market Builders"), (4.7, 1.95, "Oracle & Data"), (-4.7, 1.95, "Distribution")]
    for dx, dy, label in out_pos:
        l = int(cx + Inches(dx) - Inches(1.1)); t = int(cy + Inches(dy) - Inches(0.18))
        text(sl, label, l, t, Inches(2.2), Inches(0.36), 12, CYAN, bold=True, align=PP_ALIGN.CENTER)


def L_biz(sl, s):
    head(sl, s)
    for i, fe in enumerate(s["fees"]):
        t = Inches(3.0) + i * Inches(1.4)
        rect(sl, MX, t, Inches(6.2), Inches(1.15))
        rect(sl, MX, t, Pt(4), Inches(1.15), fill=GOLD, line=None, rad=0)
        text(sl, fe["v"], MX + Inches(0.3), t + Inches(0.18), Inches(2.2), Inches(0.8), 44, GOLD, bold=True)
        text(sl, fe["k"], MX + Inches(2.4), t + Inches(0.28), Inches(3.5), Inches(0.4), 19, WHITE, bold=True)
        text(sl, fe["s"], MX + Inches(2.4), t + Inches(0.72), Inches(3.5), Inches(0.3), 13, MUT)
    fit(sl, s["donut"], Inches(8.2), Inches(2.65), Inches(3.6), Inches(3.6))
    cx = Inches(10.0)
    text(sl, s["donut_center"][0], cx - Inches(0.8), Inches(4.2), Inches(1.6), Inches(0.4), 18, WHITE, bold=True, align=PP_ALIGN.CENTER)
    segpos = [(Inches(7.5), Inches(2.7)), (Inches(7.5), Inches(5.3)), (Inches(11.4), Inches(2.7)), (Inches(11.4), Inches(5.3))]
    for (x, y), seg in zip(segpos, s["segs"]):
        text(sl, seg["p"], x, y, Inches(1.6), Inches(0.4), 22, GOLD, bold=True)
        text(sl, seg["n"], x, y + Inches(0.38), Inches(1.7), Inches(0.3), 12, BODY)
    if s.get("note"):
        text(sl, s["note"], MX, Inches(6.7), Inches(11), Inches(0.3), 11, RGBColor(0x5C, 0x5C, 0x5C))


def L_team(sl, s):
    head(sl, s); mem = s["members"]; tw, th, gap = Inches(5.9), Inches(1.25), Inches(0.3); top = Inches(2.6)
    for i, m in enumerate(mem):
        r, c = divmod(i, 2); l = MX + c * (tw + gap); t = top + r * (th + gap)
        fit(sl, m["photo"], l, t + Inches(0.1), Inches(0.95), Inches(0.95))
        tx = l + Inches(1.15)
        text(sl, m["name"], tx, t, Inches(4.6), Inches(0.35), 17, GOLD, bold=True)
        text(sl, m["role"], tx, t + Inches(0.36), Inches(4.6), Inches(0.28), 11, MUT)
        text(sl, m["bio"], tx, t + Inches(0.66), Inches(4.6), Inches(0.6), 12, BODY, sp=1.2)
    if s.get("note"):
        text(sl, s["note"], MX, Inches(6.7), Inches(12), Inches(0.3), 11, RGBColor(0x5C, 0x5C, 0x5C))


def L_roadmap(sl, s):
    head(sl, s); ph = s["phases"]; n = len(ph); cw = int((SW - 2 * MX) / n); lt = Inches(3.0)
    rect(sl, MX, lt, SW - 2 * MX, Pt(2), fill=RGBColor(0x22, 0x22, 0x22), line=None, rad=0)
    for i, p in enumerate(ph):
        l = MX + i * cw; cxl = int(l + cw / 2 - Inches(0.13))
        done = p.get("done")
        rect(sl, cxl, lt - Inches(0.1), Inches(0.26), Inches(0.26), fill=(GOLD if done else RGBColor(0x22, 0x22, 0x22)), line=(GOLD if done else LINE), rad=0.5, shape=MSO_SHAPE.OVAL)
        text(sl, p["ph"], l, lt + Inches(0.35), cw, Inches(0.3), 10, MUT, align=PP_ALIGN.CENTER)
        text(sl, p["h"], l, lt + Inches(0.7), cw, Inches(0.4), 16, WHITE, bold=True, align=PP_ALIGN.CENTER)
        text(sl, p["t"], l + Inches(0.15), lt + Inches(1.15), cw - Inches(0.3), Inches(1.0), 12, MUT, sp=1.2, align=PP_ALIGN.CENTER)
    if s.get("closing"):
        text(sl, s["closing"], MX, Inches(5.5), SW - 2 * MX, Inches(0.8), 36, GOLD, bold=True, align=PP_ALIGN.CENTER)
    if s.get("note"):
        text(sl, s["note"], MX, Inches(6.7), SW - 2 * MX, Inches(0.3), 10, RGBColor(0x5C, 0x5C, 0x5C), align=PP_ALIGN.CENTER)


LAYOUTS = {"cover": L_cover, "investors": L_investors, "cards": L_cards, "statgrid": L_statgrid,
           "concept": L_concept, "cols": L_cols, "flow": L_flow, "spec": L_spec, "fcards": L_fcards,
           "eco": L_eco, "biz": L_biz, "team": L_team, "roadmap": L_roadmap}


def main():
    prs = Presentation(); prs.slide_width = SW; prs.slide_height = SH; blank = prs.slide_layouts[6]
    for s in SLIDES:
        sl = prs.slides.add_slide(blank); bg(sl)
        LAYOUTS[s["layout"]](sl, s)
        if s["layout"] != "cover":
            brand(sl)
    prs.save(OUT)
    print(f"✔ PPTX -> {OUT} ({os.path.getsize(OUT)//1024} KB, {len(SLIDES)} slides)")


if __name__ == "__main__":
    main()
